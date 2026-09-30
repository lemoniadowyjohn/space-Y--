from openpyxl import Workbook, load_workbook

from recon.core import reconcile_workbooks, write_report


def make_book(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["ComponentID", "Status", "Value"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_detects_mismatch_and_missing_row(tmp_path):
    left = tmp_path / "left.xlsx"
    right = tmp_path / "right.xlsx"
    make_book(
        left,
        [["A-01", "OK", 10.0], ["A-02", "OK", 20.0]],
    )
    make_book(
        right,
        [
            ["A-01", " ok ", 10.0],
            ["A-02", "NOK", 20.0],
            ["A-03", "OK", 30.0],
        ],
    )
    result = reconcile_workbooks(left, right)
    assert result.matched_rows == 2
    assert len(result.differences) == 2
    assert {d.kind for d in result.differences} == {
        "value_mismatch",
        "missing_left",
    }


def test_duplicate_key_is_rejected(tmp_path):
    left = tmp_path / "left.xlsx"
    right = tmp_path / "right.xlsx"
    make_book(
        left,
        [["A-01", "OK", 1], ["A-01", "OK", 1]],
    )
    make_book(right, [["A-01", "OK", 1]])

    try:
        reconcile_workbooks(left, right)
    except ValueError as exc:
        assert "duplicate key" in str(exc)
    else:
        raise AssertionError("duplicate key should fail")


def test_report_is_written(tmp_path):
    left = tmp_path / "left.xlsx"
    right = tmp_path / "right.xlsx"
    report = tmp_path / "report.xlsx"

    make_book(left, [["A-01", "OK", 1]])
    make_book(right, [["A-01", "NOK", 1]])

    result = reconcile_workbooks(left, right)
    write_report(result, report)

    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B2"].value is False
    assert wb["Summary"]["B6"].value == 1

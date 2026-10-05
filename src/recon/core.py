from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class Difference:
    key: str
    field: str
    left: Any
    right: Any
    kind: str


@dataclass(frozen=True)
class ReconciliationResult:
    differences: tuple[Difference, ...]
    matched_rows: int
    left_rows: int
    right_rows: int

    @property
    def passed(self) -> bool:
        return not self.differences


def _normalize(value: Any) -> Any:
    if isinstance(value, str):
        return " ".join(value.strip().split()).casefold()
    if isinstance(value, float):
        return round(value, 6)
    return value


def _read_table(
    path: str | Path,
    sheet: str,
    key_column: str,
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers: list[str] = []
    for cell in ws[1]:
        if not isinstance(cell.value, str):
            raise ValueError("header row must contain non-empty string column names")
        headers.append(cell.value)
    if key_column not in headers:
        raise ValueError(f"missing key column: {key_column}")

    rows: dict[str, dict[str, Any]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        row: dict[str, Any] = {
            header: value
            for header, value in zip(headers, values, strict=True)
        }
        key = str(row[key_column]).strip()
        if key in rows:
            raise ValueError(f"duplicate key: {key}")
        rows[key] = row
    return rows, headers


def reconcile_workbooks(
    left_path: str | Path,
    right_path: str | Path,
    *,
    sheet: str = "Data",
    key_column: str = "ComponentID",
) -> ReconciliationResult:
    left, left_headers = _read_table(left_path, sheet, key_column)
    right, right_headers = _read_table(right_path, sheet, key_column)
    fields = [
        header
        for header in left_headers
        if header != key_column and header in right_headers
    ]

    differences: list[Difference] = []
    matched = 0

    for key in sorted(set(left) | set(right)):
        if key not in left:
            differences.append(
                Difference(key, key_column, None, key, "missing_left")
            )
            continue
        if key not in right:
            differences.append(
                Difference(key, key_column, key, None, "missing_right")
            )
            continue

        matched += 1
        for field in fields:
            left_value = left[key].get(field)
            right_value = right[key].get(field)
            if _normalize(left_value) != _normalize(right_value):
                differences.append(
                    Difference(
                        key,
                        field,
                        left_value,
                        right_value,
                        "value_mismatch",
                    )
                )

    return ReconciliationResult(
        tuple(differences),
        matched,
        len(left),
        len(right),
    )


def write_report(
    result: ReconciliationResult,
    path: str | Path,
) -> None:
    wb = Workbook()

    detail = wb.active
    if detail is None:
        raise RuntimeError("workbook did not create an active worksheet")
    detail.title = "Reconciliation"
    detail.append(["Key", "Field", "Left", "Right", "Kind"])
    for diff in result.differences:
        detail.append(
            [diff.key, diff.field, diff.left, diff.right, diff.kind]
        )

    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Passed", result.passed])
    summary.append(["Matched rows", result.matched_rows])
    summary.append(["Left rows", result.left_rows])
    summary.append(["Right rows", result.right_rows])
    summary.append(["Differences", len(result.differences)])

    wb.save(path)
